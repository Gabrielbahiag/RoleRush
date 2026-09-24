from datetime import date
from pathlib import Path

import pytest
from docx import Document

from monitor.curriculo.documento import gerar_docx, nome_de_arquivo
from monitor.curriculo.extracao import RequisitosVaga
from monitor.curriculo.mestre import carregar_curriculo
from monitor.curriculo.selecao import selecionar
from monitor.curriculo.skills import carregar_skills

RAIZ = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def dicionario():
    return carregar_skills(RAIZ / "config" / "skills.yaml")


@pytest.fixture(scope="module")
def mestre():
    return carregar_curriculo(RAIZ / "curriculo_mestre.example.yaml")


def _adaptado(mestre, dicionario, idioma="pt"):
    requisitos = RequisitosVaga(obrigatorios=["Python", "SQL"], desejaveis=["Pandas"], idioma=idioma)
    return selecionar(mestre, requisitos, dicionario, idioma=idioma)


def _paragrafos(caminho: Path) -> list[str]:
    return [p.text.strip() for p in Document(str(caminho)).paragraphs if p.text.strip()]


def test_arquivo_gerado_abre_e_tem_as_secoes_na_ordem(tmp_path, mestre, dicionario):
    destino = tmp_path / "cv.docx"

    gerar_docx(_adaptado(mestre, dicionario), destino)

    textos = _paragrafos(destino)
    esperadas = ["Resumo", "Experiência Profissional", "Projetos", "Formação", "Competências"]
    posicoes = [textos.index(secao) for secao in esperadas]
    assert posicoes == sorted(posicoes)


def test_secoes_saem_no_idioma_pedido(tmp_path, mestre, dicionario):
    destino = tmp_path / "cv.docx"

    gerar_docx(_adaptado(mestre, dicionario, idioma="en"), destino)

    textos = _paragrafos(destino)
    assert "Professional Experience" in textos
    assert "Experiência Profissional" not in textos


def test_documento_traz_nome_contato_e_conteudo_selecionado(tmp_path, mestre, dicionario):
    destino = tmp_path / "cv.docx"
    adaptado = _adaptado(mestre, dicionario)

    gerar_docx(adaptado, destino)

    conteudo = "\n".join(_paragrafos(destino))
    assert mestre.dados.nome in conteudo
    assert mestre.dados.contato.email in conteudo
    primeiro_bullet = adaptado.experiencias[0].bullets[0]
    assert primeiro_bullet.texto_para("pt") in conteudo


def test_estrutura_e_ats_sem_tabela_e_sem_imagem(tmp_path, mestre, dicionario):
    destino = tmp_path / "cv.docx"

    gerar_docx(_adaptado(mestre, dicionario), destino)

    documento = Document(str(destino))
    assert documento.tables == []
    assert len(documento.inline_shapes) == 0


def test_gerar_docx_devolve_o_caminho_e_cria_o_diretorio(tmp_path, mestre, dicionario):
    destino = tmp_path / "saida" / "cv.docx"

    resultado = gerar_docx(_adaptado(mestre, dicionario), destino)

    assert resultado == destino
    assert destino.exists()


@pytest.mark.parametrize(
    "empresa,cargo,esperado",
    [
        ("Acme", "Dev Python", "acme-dev-python-2026-09-24.docx"),
        # regra: toda sequência de não-alfanumérico vira UM hífen — por isso
        # "S/A" sai como "s-a". É feio, mas é previsível.
        ("Acme S/A", "Desenvolvedor Sênior", "acme-s-a-desenvolvedor-senior-2026-09-24.docx"),
        ("  Órbita  Tech  ", "Analista de Dados", "orbita-tech-analista-de-dados-2026-09-24.docx"),
        ("", "", "curriculo-2026-09-24.docx"),
    ],
)
def test_nome_de_arquivo_e_sanitizado(empresa, cargo, esperado):
    assert nome_de_arquivo(empresa, cargo, date(2026, 9, 24)) == esperado
