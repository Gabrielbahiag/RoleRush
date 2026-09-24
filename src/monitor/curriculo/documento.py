from __future__ import annotations

import re
import unicodedata
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

from monitor.curriculo.normalizacao import Idioma
from monitor.curriculo.selecao import CurriculoAdaptado

# ATS lê texto corrido: uma coluna, sem tabela, sem imagem, sem caixa de texto,
# títulos de seção previsíveis e fonte comum.
_FONTE = "Calibri"
_TAMANHO_BASE = Pt(11)

_TITULOS: dict[Idioma, dict[str, str]] = {
    "pt": {
        "resumo": "Resumo",
        "experiencia": "Experiência Profissional",
        "projetos": "Projetos",
        "formacao": "Formação",
        "skills": "Competências",
        "idiomas": "Idiomas",
        "certificacoes": "Certificações",
    },
    "en": {
        "resumo": "Summary",
        "experiencia": "Professional Experience",
        "projetos": "Projects",
        "formacao": "Education",
        "skills": "Skills",
        "idiomas": "Languages",
        "certificacoes": "Certifications",
    },
}


def gerar_docx(
    adaptado: CurriculoAdaptado,
    caminho: str | Path,
    reescritas: dict[str, str] | None = None,
) -> Path:
    """Escreve o .docx. `reescritas` (id do bullet -> texto) entra no lugar do
    original só para o que passou pelo guardrail; o mestre nunca é alterado."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    reescritas = reescritas or {}

    idioma = adaptado.idioma
    titulos = _TITULOS[idioma]
    mestre = adaptado.mestre

    documento = Document()
    _preparar_estilo(documento)

    _paragrafo(documento, mestre.dados.nome, negrito=True, tamanho=Pt(16))
    _paragrafo(documento, " · ".join(_contato(mestre)))
    links = _links(mestre)
    if links:
        _paragrafo(documento, " · ".join(links))

    _secao(documento, titulos["resumo"])
    _paragrafo(documento, mestre.resumo.para(idioma, "resumo"))

    if adaptado.experiencias:
        _secao(documento, titulos["experiencia"])
        for selecionada in adaptado.experiencias:
            experiencia = selecionada.experiencia
            _paragrafo(
                documento,
                f"{experiencia.cargo} — {experiencia.empresa} ({experiencia.periodo})",
                negrito=True,
            )
            for bullet in selecionada.bullets:
                _item(documento, reescritas.get(bullet.id) or bullet.texto_para(idioma))

    if adaptado.projetos:
        _secao(documento, titulos["projetos"])
        for selecionado in adaptado.projetos:
            projeto = selecionado.projeto
            cabecalho = projeto.nome
            if projeto.stack:
                cabecalho += f" — {', '.join(projeto.stack)}"
            _paragrafo(documento, cabecalho, negrito=True)
            _paragrafo(documento, projeto.descricao.para(idioma, f"projeto '{projeto.nome}'"))
            if projeto.link:
                _paragrafo(documento, projeto.link)
            for bullet in selecionado.bullets:
                _item(documento, reescritas.get(bullet.id) or bullet.texto_para(idioma))

    if mestre.formacao:
        _secao(documento, titulos["formacao"])
        for formacao in mestre.formacao:
            linha = f"{formacao.curso} — {formacao.instituicao}"
            if formacao.status:
                linha += f" ({formacao.status})"
            _paragrafo(documento, linha)

    if adaptado.skills:
        _secao(documento, titulos["skills"])
        _paragrafo(documento, ", ".join(adaptado.skills))

    if mestre.idiomas:
        _secao(documento, titulos["idiomas"])
        _paragrafo(
            documento,
            ", ".join(
                f"{idioma_falado.nome} ({idioma_falado.nivel})" if idioma_falado.nivel else idioma_falado.nome
                for idioma_falado in mestre.idiomas
            ),
        )

    if mestre.certificacoes:
        _secao(documento, titulos["certificacoes"])
        for certificacao in mestre.certificacoes:
            partes = [certificacao.nome]
            if certificacao.emissor:
                partes.append(certificacao.emissor)
            if certificacao.ano:
                partes.append(str(certificacao.ano))
            _paragrafo(documento, " — ".join(partes))

    documento.save(str(caminho))
    return caminho


def nome_de_arquivo(empresa: str, cargo: str, dia: date | None = None) -> str:
    dia = dia or date.today()
    partes = [parte for parte in (_sanitizar(empresa), _sanitizar(cargo)) if parte]
    if not partes:
        partes = ["curriculo"]
    return "-".join([*partes, dia.isoformat()]) + ".docx"


def _sanitizar(texto: str) -> str:
    # acento vira ascii, o resto que não for letra/número vira hífen: nome de
    # arquivo tem que sobreviver a e-mail, upload e sistema de arquivos alheio.
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"-{2,}", "-", re.sub(r"[^a-z0-9]+", "-", sem_acento.lower())).strip("-")


def _preparar_estilo(documento: Document) -> None:
    estilo = documento.styles["Normal"]
    estilo.font.name = _FONTE
    estilo.font.size = _TAMANHO_BASE


def _contato(mestre) -> list[str]:
    contato = mestre.dados.contato
    return [valor for valor in (contato.email, contato.telefone, contato.cidade) if valor]


def _links(mestre) -> list[str]:
    links = mestre.dados.links
    return [valor for valor in (links.github, links.linkedin, links.portfolio) if valor]


def _secao(documento: Document, titulo: str) -> None:
    _paragrafo(documento, titulo, negrito=True, tamanho=Pt(13))


def _paragrafo(documento: Document, texto: str, negrito: bool = False, tamanho: Pt | None = None):
    paragrafo = documento.add_paragraph()
    corrida = paragrafo.add_run(texto)
    corrida.bold = negrito
    if tamanho is not None:
        corrida.font.size = tamanho
    return paragrafo


def _item(documento: Document, texto: str) -> None:
    documento.add_paragraph(texto, style="List Bullet")
